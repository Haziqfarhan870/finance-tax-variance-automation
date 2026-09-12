from __future__ import annotations
from pathlib import Path
import pandas as pd
from openpyxl import Workbook
from openpyxl.styles import Font, PatternFill, Alignment, Border, Side
from openpyxl.utils import get_column_letter
from openpyxl.worksheet.datavalidation import DataValidation

BORDER = Border(left=Side(style="thin", color="D9D9D9"), right=Side(style="thin", color="D9D9D9"), top=Side(style="thin", color="D9D9D9"), bottom=Side(style="thin", color="D9D9D9"))


def _quarterly_sheet(wb, name: str, df: pd.DataFrame, quarters: list[str], comments: dict[str, tuple[str, str]]) -> tuple[str, dict[str, tuple[int,int]]]:
    ws = wb.create_sheet(name)
    headers = ["Company Code","Level 1","Level 2","Account Name","Account"] + quarters + ["Delta QoQ","% Change QoQ","Materiality","YTD","Trend","Review Status","Variance Comment","Account Insight"]
    for c,h in enumerate(headers,1):
        cell=ws.cell(1,c,h); cell.font=Font(bold=True,color="FFFFFF"); cell.fill=PatternFill("solid",fgColor="4472C4"); cell.alignment=Alignment(horizontal="center"); cell.border=BORDER
    ranges={}; rownum=2
    for comp in sorted(df["Company Code"].astype(str).unique()):
        start=rownum
        for _,r in df[df["Company Code"].astype(str)==comp].iterrows():
            if str(r["Level 1"])=="UNMAPPED":
                continue
            key=f"{r['Company Code']}_{r['Account']}"; comment, insight = comments.get(key,("",""))
            vals=[r["Company Code"],r["Level 1"],r["Level 2"],r["Account Name"],r["Account"]] + [r[q] for q in quarters] + [r["Delta QoQ"],r["% Change QoQ"],r["Materiality"],r["YTD"],r["Trend"],"",comment,insight]
            for c,v in enumerate(vals,1):
                cell=ws.cell(rownum,c,v); cell.border=BORDER; cell.alignment=Alignment(vertical="top",wrap_text=c>=len(headers)-1)
                if headers[c-1] in quarters+["Delta QoQ","YTD"]: cell.number_format='#,##0;[Red](#,##0);-'
                if headers[c-1]=="% Change QoQ": cell.number_format='0.0%'; cell.value=float(v)/100 if v not in (None,"") else 0
            if r["Materiality"]=="❗":
                ws.cell(rownum, headers.index("Materiality")+1).fill=PatternFill("solid",fgColor="FFC7CE")
            rownum+=1
        ranges[comp]=(start,rownum-1)
    review_col=headers.index("Review Status")+1
    review_letter=get_column_letter(review_col)
    dv=DataValidation(type="list",formula1='"Not Started,In Progress,Reviewed"',allow_blank=True); ws.add_data_validation(dv); dv.add(f"{review_letter}2:{review_letter}{max(2,rownum-1)}")
    ws.freeze_panes="A2"; ws.auto_filter.ref=f"A1:{get_column_letter(len(headers))}{max(2,rownum-1)}"
    widths=[14,20,20,28,12]+[14]*len(quarters)+[14,14,12,14,10,16,55,55]
    for i,w in enumerate(widths,1): ws.column_dimensions[get_column_letter(i)].width=w
    return review_letter,ranges


def _summary_sheet(wb, name: str, df: pd.DataFrame, quarterly_name: str, review_col: str, ranges: dict[str,tuple[int,int]]) -> None:
    ws=wb.create_sheet(name)
    headers=["Company Code","Latest Quarter Total","QoQ Movement","Material Variances","Total Accounts","Review Completion %"]
    for c,h in enumerate(headers,1):
        cell=ws.cell(4,c,h); cell.font=Font(bold=True); cell.fill=PatternFill("solid",fgColor="D9E1F2"); cell.border=BORDER; cell.alignment=Alignment(horizontal="center")
    latest_candidates=[c for c in df.columns if isinstance(c,str) and "Q" in c and c[0].isdigit()]
    latest=latest_candidates[-1]
    row=5
    for comp in sorted(df["Company Code"].astype(str).unique()):
        part=df[df["Company Code"].astype(str)==comp]
        total_accounts=int((part["Level 1"]!="UNMAPPED").sum())
        vals=[comp,float(part[latest].sum()),float(part["Delta QoQ"].sum()),int((part["Materiality"]=="❗").sum()),total_accounts,None]
        for c,v in enumerate(vals,1): ws.cell(row,c,v).border=BORDER
        start,end=ranges.get(comp,(None,None))
        if start and end and end>=start and total_accounts:
            ws.cell(row,6,f'=IFERROR(COUNTIF(\'{quarterly_name}\'!{review_col}{start}:{review_col}{end},"Reviewed")/{total_accounts},0)')
        else: ws.cell(row,6,0)
        ws.cell(row,6).number_format='0%'
        for c in [2,3]: ws.cell(row,c).number_format='#,##0;[Red](#,##0);-'
        row+=1
    ws["A1"]=name.upper(); ws["A1"].font=Font(size=14,bold=True,color="FFFFFF"); ws["A1"].fill=PatternFill("solid",fgColor="305496"); ws.merge_cells("A1:F1")
    ws["A2"]=f"Latest quarter: {latest} | Review completion updates from the detailed sheet"; ws.merge_cells("A2:F2")
    for i,w in enumerate([16,22,20,20,16,22],1): ws.column_dimensions[get_column_letter(i)].width=w
    ws.freeze_panes="A5"


def write_report(pnl_df: pd.DataFrame, pnl_q: list[str], pnl_comments: dict, bs_df: pd.DataFrame, bs_q: list[str], bs_comments: dict, output: Path) -> None:
    wb=Workbook(); wb.remove(wb.active)
    if not pnl_df.empty:
        review,ranges=_quarterly_sheet(wb,"PNL Quarterly Analysis",pnl_df,pnl_q,pnl_comments); _summary_sheet(wb,"PNL Summary",pnl_df,"PNL Quarterly Analysis",review,ranges)
    if not bs_df.empty:
        review,ranges=_quarterly_sheet(wb,"BS Quarterly Analysis",bs_df,bs_q,bs_comments); _summary_sheet(wb,"BS Summary",bs_df,"BS Quarterly Analysis",review,ranges)
    order=[s for s in ["PNL Summary","PNL Quarterly Analysis","BS Summary","BS Quarterly Analysis"] if s in wb.sheetnames]
    wb._sheets=[wb[s] for s in order]
    output.parent.mkdir(parents=True,exist_ok=True); wb.save(output)
