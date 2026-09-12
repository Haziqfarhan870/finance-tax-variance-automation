from pathlib import Path
import subprocess,sys
from openpyxl import load_workbook
from src.pipeline import run

def test_end_to_end(tmp_path):
    root=Path(__file__).resolve().parents[1]
    # create demo data in the repo fixture area, then run real pipeline there
    subprocess.check_call([sys.executable,str(root/'scripts'/'generate_synthetic_data.py')])
    out=run(root)
    assert out.exists()
    wb=load_workbook(out,read_only=False,data_only=False)
    assert wb.sheetnames==['PNL Summary','PNL Quarterly Analysis','BS Summary','BS Quarterly Analysis']
    assert wb['PNL Quarterly Analysis'].data_validations.count>=1
